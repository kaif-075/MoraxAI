import tkinter as tk
from tkinter import ttk, messagebox
import threading
import time
import re
import os
from datetime import datetime

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

SBTET_URL = "https://sbtet.ap.gov.in/APSBTET/gradeWiseResults.do"

# Results will be saved in the same folder as this Python file
OUTPUT_FOLDER = os.path.join(os.path.dirname(os.path.abspath(__file__)))

STOP_SCRAPING = False


# ============================================================
# HELPER FUNCTIONS
# ============================================================

def clean_text(text):
    """Clean unnecessary spaces/newlines."""
    if text is None:
        return ""

    return " ".join(text.replace("\xa0", " ").split()).strip()


def get_numeric_value(text):
    """
    Keep marks clean.
    Example:
        '33' -> 33
        '11' -> 11
        '-'  -> ''
    """

    text = clean_text(text)

    if not text:
        return ""

    if text in ["-", "--", "NA", "N/A"]:
        return ""

    # Integer
    if re.fullmatch(r"-?\d+", text):
        return int(text)

    # Decimal
    if re.fullmatch(r"-?\d+\.\d+", text):
        return float(text)

    return text


def create_driver():

    options = webdriver.EdgeOptions()

    # Run Edge without opening a visible browser
    options.add_argument("--headless=new")
    options.add_argument("--disable-gpu")
    options.add_argument("--window-size=1920,1080")
    options.add_argument("--log-level=3")

    # Selenium Manager automatically finds/downloads
    # the correct EdgeDriver.
    #
    # This is better than:
    # Service(r"C:\WebDriver\msedgedriver.exe")
    #
    # because EdgeDriver versions change.
    driver = webdriver.Edge(options=options)

    return driver


# ============================================================
# SELECT MONTH/YEAR
# ============================================================

def select_month_year(driver, month_year):

    selects = driver.find_elements(By.TAG_NAME, "select")

    if len(selects) < 2:
        raise Exception("Could not find Month/Year and Semester dropdowns.")

    month_dropdown = None

    # Find the dropdown which contains the requested month/year
    for select_element in selects:

        try:
            options = Select(select_element).options

            option_texts = [
                clean_text(option.text).upper()
                for option in options
            ]

            if month_year.upper() in option_texts:
                month_dropdown = select_element
                break

        except Exception:
            continue

    if month_dropdown is None:

        # Print available options for debugging
        available = []

        for select_element in selects:
            try:
                options = Select(select_element).options

                available.extend([
                    clean_text(option.text)
                    for option in options
                ])

            except:
                pass

        raise Exception(
            f"Month/Year '{month_year}' was not found.\n"
            f"Available options:\n{available}"
        )

    Select(month_dropdown).select_by_visible_text(month_year)


# ============================================================
# SELECT SEMESTER
# ============================================================

def select_semester(driver, semester):

    selects = driver.find_elements(By.TAG_NAME, "select")

    semester_dropdown = None

    for select_element in selects:

        try:

            options = Select(select_element).options

            option_texts = [
                clean_text(option.text).upper()
                for option in options
            ]

            if semester.upper() in option_texts:

                semester_dropdown = select_element
                break

        except:
            continue

    if semester_dropdown is None:
        raise Exception(
            f"Semester '{semester}' was not found."
        )

    Select(semester_dropdown).select_by_visible_text(semester)


# ============================================================
# EXTRACT STUDENT DETAILS
# ============================================================

def extract_student_details(driver):

    data = {
        "student_name": "",
        "branch": "",
        "papers": {},
        "grand_total": "",
        "gpa": "",
        "result": ""
    }

    # --------------------------------------------------------
    # STUDENT NAME
    # --------------------------------------------------------

    try:

        data["student_name"] = clean_text(
            driver.find_element(
                By.XPATH,
                "//th[normalize-space()='Name']/following-sibling::td[1]"
            ).text
        )

    except:
        pass


    # --------------------------------------------------------
    # BRANCH
    # --------------------------------------------------------

    try:

        data["branch"] = clean_text(
            driver.find_element(
                By.XPATH,
                "//th[normalize-space()='Branch Name']/following-sibling::td[1]"
            ).text
        )

    except:
        pass


    # --------------------------------------------------------
    # GRAND TOTAL
    # --------------------------------------------------------

    try:

        data["grand_total"] = get_numeric_value(
            driver.find_element(
                By.XPATH,
                "//th[normalize-space()='Grand Total']/following-sibling::td[1]"
            ).text
        )

    except:
        pass


    # --------------------------------------------------------
    # GPA
    # --------------------------------------------------------

    try:

        data["gpa"] = clean_text(
            driver.find_element(
                By.XPATH,
                "//th[normalize-space()='GPA']/following-sibling::td[1]"
            ).text
        )

    except:
        pass


    # --------------------------------------------------------
    # RESULT
    # --------------------------------------------------------

    try:

        data["result"] = clean_text(
            driver.find_element(
                By.XPATH,
                "//th[normalize-space()='Result']/following-sibling::td[1]"
            ).text
        )

    except:
        pass


    # ========================================================
    # SUBJECT / PAPER MARKS
    # ========================================================

    # The result table contains rows like:
    #
    # Paper | External | Internal | Total | Grade Points |
    # Credits Earned | Grade | SUB.Result
    #
    # Example:
    #
    # 401 | 33 | 11 | 44 | 5 | 3.0 | C | P
    #
    # We only save:
    #
    # 401 External = 33
    # 401 Internal = 11
    # 401 Total    = 44

    try:

        tables = driver.find_elements(
            By.CSS_SELECTOR,
            "#printDiv table"
        )

        result_table = None

        # Find table containing "Paper"
        for table in tables:

            table_text = clean_text(table.text).upper()

            if "PAPER" in table_text and "EXTERNAL" in table_text:

                result_table = table
                break

        if result_table is None:

            # Fallback
            tables = driver.find_elements(By.TAG_NAME, "table")

            for table in tables:

                table_text = clean_text(table.text).upper()

                if "PAPER" in table_text and "EXTERNAL" in table_text:

                    result_table = table
                    break


        if result_table is not None:

            rows = result_table.find_elements(
                By.TAG_NAME,
                "tr"
            )

            for row in rows:

                cells = row.find_elements(
                    By.TAG_NAME,
                    "th"
                )

                cells += row.find_elements(
                    By.TAG_NAME,
                    "td"
                )

                if len(cells) < 4:
                    continue

                values = [
                    clean_text(cell.text)
                    for cell in cells
                ]

                # First column must look like a paper number.
                #
                # Examples:
                # 401
                # 402
                # 403
                #
                if not re.fullmatch(
                    r"\d{3,4}",
                    values[0]
                ):
                    continue

                paper_no = values[0]

                external = get_numeric_value(values[1])
                internal = get_numeric_value(values[2])
                total = get_numeric_value(values[3])

                data["papers"][paper_no] = {
                    "external": external,
                    "internal": internal,
                    "total": total
                }

    except Exception as e:

        print(
            "Subject extraction error:",
            str(e)
        )


    return data


# ============================================================
# FETCH ONE RESULT
# ============================================================

def fetch_result(
    driver,
    wait,
    pin,
    month_year,
    semester
):

    try:

        driver.get(SBTET_URL)

        # ----------------------------------------------------
        # PIN
        # ----------------------------------------------------

        # PIN
        pin_input = wait.until(
            EC.presence_of_element_located((By.ID, "aadhar1"))
            )
        pin_input.clear()
        pin_input.send_keys(pin)

        # Month / Year
        month_dropdown = wait.until(
            EC.presence_of_element_located((By.ID, "feesmonthyear"))
        )

        Select(month_dropdown).select_by_visible_text(month_year)

        # Semester
        semester_dropdown = wait.until(
            EC.presence_of_element_located((By.ID, "grade2"))
        )

        Select(semester_dropdown).select_by_visible_text(semester)

        # Submit
        submit_button = wait.until(
            EC.element_to_be_clickable(
                (By.CSS_SELECTOR, "input[type='button'][value='SUBMIT']")
            )
        )

        submit_button.click()


        # ----------------------------------------------------
        # WAIT FOR RESULT PAGE
        # ----------------------------------------------------

        time.sleep(2)


        # ----------------------------------------------------
        # CHECK PAGE TEXT
        # ----------------------------------------------------

        page_text = driver.page_source.upper()

        error_messages = [
            "NOT REGISTERED",
            "REGISTERED BUT NOT ELIGIBLE",
            "NO RECORDS",
            "NO RECORD FOUND",
            "NO RESULTS FOUND"
        ]

        for error_message in error_messages:

            if error_message in page_text:

                return {
                    "success": False,
                    "error": "No result found."
                }


        # ----------------------------------------------------
        # WAIT FOR RESULT TABLE
        # ----------------------------------------------------

        try:

            wait.until(
                EC.presence_of_element_located(
                    (
                        By.XPATH,
                        "//th[normalize-space()='Name']"
                    )
                )
            )

        except:

            pass


        # ----------------------------------------------------
        # EXTRACT DATA
        # ----------------------------------------------------

        result = extract_student_details(
            driver
        )


        # ----------------------------------------------------
        # VERIFY THAT RESULT WAS ACTUALLY FOUND
        # ----------------------------------------------------

        if not result["student_name"]:

            return {
                "success": False,
                "error": "Student result could not be read."
            }


        result["success"] = True

        return result


    except Exception as e:

        return {
            "success": False,
            "error": str(e)
        }


# ============================================================
# SAVE EXCEL FILE
# ============================================================

def save_excel(results, output_file):

    # --------------------------------------------------------
    # SORT BY PIN
    # --------------------------------------------------------

    results.sort(
        key=lambda x: int(
            re.search(
                r"(\d+)$",
                x["pin"]
            ).group(1)
        )
    )


    # --------------------------------------------------------
    # COLLECT ALL PAPER NUMBERS
    # --------------------------------------------------------

    paper_numbers = set()

    for student in results:

        for paper in student["papers"]:

            paper_numbers.add(paper)


    # Increasing paper order
    paper_numbers = sorted(
        paper_numbers,
        key=lambda x: int(x)
    )


    # --------------------------------------------------------
    # CREATE WORKBOOK
    # --------------------------------------------------------

    workbook = Workbook()

    sheet = workbook.active

    sheet.title = "SBTET Results"


    # --------------------------------------------------------
    # HEADERS
    # --------------------------------------------------------

    headers = [
        "PIN",
        "Student Name",
        "Branch"
    ]


    # Add:

    # 401 External
    # 401 Internal
    # 401 Total
    #
    # 402 External
    # 402 Internal
    # 402 Total
    #
    # etc.

    for paper in paper_numbers:

        headers.append(
            f"{paper} External"
        )

        headers.append(
            f"{paper} Internal"
        )

        headers.append(
            f"{paper} Total"
        )


    headers.extend([
        "Grand Total",
        "GPA",
        "Result"
    ])


    # --------------------------------------------------------
    # WRITE HEADERS
    # --------------------------------------------------------

    for column, header in enumerate(
        headers,
        start=1
    ):

        cell = sheet.cell(
            row=1,
            column=column,
            value=header
        )

        cell.font = Font(
            bold=True
        )

        cell.alignment = Alignment(
            horizontal="center",
            vertical="center"
        )


    # --------------------------------------------------------
    # WRITE STUDENTS
    # --------------------------------------------------------

    excel_row = 2

    for student in results:

        row = [
            student["pin"],
            student["student_name"],
            student["branch"]
        ]


        # ----------------------------------------------------
        # SUBJECT MARKS
        # ----------------------------------------------------

        for paper in paper_numbers:

            paper_data = student["papers"].get(
                paper,
                {}
            )

            row.append(
                paper_data.get(
                    "external",
                    ""
                )
            )

            row.append(
                paper_data.get(
                    "internal",
                    ""
                )
            )

            row.append(
                paper_data.get(
                    "total",
                    ""
                )
            )


        # ----------------------------------------------------
        # FINAL DETAILS
        # ----------------------------------------------------

        row.extend([
            student["grand_total"],
            student["gpa"],
            student["result"]
        ])


        # ----------------------------------------------------
        # WRITE ROW
        # ----------------------------------------------------

        for column, value in enumerate(
            row,
            start=1
        ):

            cell = sheet.cell(
                row=excel_row,
                column=column,
                value=value
            )

            cell.alignment = Alignment(
                vertical="center"
            )


        excel_row += 1


    # --------------------------------------------------------
    # FREEZE HEADER
    # --------------------------------------------------------

    sheet.freeze_panes = "A2"


    # --------------------------------------------------------
    # AUTO COLUMN WIDTH
    # --------------------------------------------------------

    for column_cells in sheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            try:

                value_length = len(
                    str(cell.value)
                )

                if value_length > max_length:
                    max_length = value_length

            except:
                pass


        # Limit width
        adjusted_width = min(
            max_length + 2,
            35
        )

        sheet.column_dimensions[
            column_letter
        ].width = adjusted_width


    # --------------------------------------------------------
    # HEADER HEIGHT
    # --------------------------------------------------------

    sheet.row_dimensions[1].height = 30


    # --------------------------------------------------------
    # SAVE
    # --------------------------------------------------------

    workbook.save(
        output_file
    )


# ============================================================
# GUI APPLICATION
# ============================================================

class SBTETDownloader:

    def __init__(self, root):

        self.root = root

        self.root.title(
            "SBTET Result Downloader"
        )

        self.root.geometry(
            "900x720"
        )

        self.root.resizable(
            False,
            False
        )


        self.driver = None


        # ====================================================
        # TITLE
        # ====================================================

        title = tk.Label(
            root,
            text="SBTET Result Downloader",
            font=("Arial", 22, "bold")
        )

        title.pack(
            pady=(20, 5)
        )


        subtitle = tk.Label(
            root,
            text="Diploma C23 Result Automation",
            font=("Arial", 12)
        )

        subtitle.pack(
            pady=(0, 15)
        )


        # ====================================================
        # SETTINGS FRAME
        # ====================================================

        settings_frame = ttk.LabelFrame(
            root,
            text="Result Settings",
            padding=15
        )

        settings_frame.pack(
            padx=25,
            fill="x"
        )


        # ====================================================
        # PIN PREFIX
        # ====================================================

        ttk.Label(
            settings_frame,
            text="PIN Prefix:"
        ).grid(
            row=0,
            column=0,
            sticky="w",
            padx=5,
            pady=7
        )


        self.prefix_entry = ttk.Entry(
            settings_frame,
            width=80
        )

        self.prefix_entry.insert(
            0,
            "24262-CM-"
        )

        self.prefix_entry.grid(
            row=0,
            column=1,
            padx=5,
            pady=7
        )


        # ====================================================
        # START PIN
        # ====================================================

        ttk.Label(
            settings_frame,
            text="Start PIN:"
        ).grid(
            row=1,
            column=0,
            sticky="w",
            padx=5,
            pady=7
        )


        self.start_entry = ttk.Entry(
            settings_frame,
            width=80
        )

        self.start_entry.insert(
            0,
            "001"
        )

        self.start_entry.grid(
            row=1,
            column=1,
            padx=5,
            pady=7
        )


        # ====================================================
        # END PIN
        # ====================================================

        ttk.Label(
            settings_frame,
            text="End PIN:"
        ).grid(
            row=2,
            column=0,
            sticky="w",
            padx=5,
            pady=7
        )


        self.end_entry = ttk.Entry(
            settings_frame,
            width=80
        )

        self.end_entry.insert(
            0,
            "033"
        )

        self.end_entry.grid(
            row=2,
            column=1,
            padx=5,
            pady=7
        )


        # ====================================================
        # MONTH / YEAR
        # ====================================================

        ttk.Label(
            settings_frame,
            text="Month / Year:"
        ).grid(
            row=3,
            column=0,
            sticky="w",
            padx=5,
            pady=7
        )


        self.month_entry = ttk.Entry(
            settings_frame,
            width=80
        )

        self.month_entry.insert(
            0,
            "MA26"
        )

        self.month_entry.grid(
            row=3,
            column=1,
            padx=5,
            pady=7
        )


        # ====================================================
        # SEMESTER
        # ====================================================

        ttk.Label(
            settings_frame,
            text="Semester:"
        ).grid(
            row=4,
            column=0,
            sticky="w",
            padx=5,
            pady=7
        )


        self.semester_combo = ttk.Combobox(
            settings_frame,
            width=77,
            state="normal"
        )

        self.semester_combo["values"] = (
            "1SEM",
            "2SEM",
            "3SEM",
            "4SEM",
            "5SEM",
            "6SEM"
        )

        self.semester_combo.set(
            "4SEM"
        )

        self.semester_combo.grid(
            row=4,
            column=1,
            padx=5,
            pady=7
        )


        # ====================================================
        # BUTTONS
        # ====================================================

        button_frame = tk.Frame(
            root
        )

        button_frame.pack(
            pady=20
        )


        self.start_button = ttk.Button(
            button_frame,
            text="START SCRAPING",
            command=self.start_scraping
        )

        self.start_button.grid(
            row=0,
            column=0,
            padx=10
        )


        self.stop_button = ttk.Button(
            button_frame,
            text="STOP",
            command=self.stop_scraping
        )

        self.stop_button.grid(
            row=0,
            column=1,
            padx=10
        )


        # ====================================================
        # STATUS
        # ====================================================

        self.status_label = tk.Label(
            root,
            text="Ready",
            font=("Arial", 12, "bold")
        )

        self.status_label.pack(
            pady=5
        )


        # ====================================================
        # PROGRESS BAR
        # ====================================================

        self.progress = ttk.Progressbar(
            root,
            orient="horizontal",
            length=810,
            mode="determinate"
        )

        self.progress.pack(
            pady=5
        )


        # ====================================================
        # LOG FRAME
        # ====================================================

        log_frame = ttk.LabelFrame(
            root,
            text="Progress Log",
            padding=5
        )

        log_frame.pack(
            padx=25,
            pady=10,
            fill="both",
            expand=True
        )


        self.log_text = tk.Text(
            log_frame,
            height=12,
            width=100
        )

        self.log_text.pack(
            fill="both",
            expand=True
        )


    # ========================================================
    # LOG
    # ========================================================

    def log(self, message):

        timestamp = datetime.now().strftime(
            "%H:%M:%S"
        )

        text = f"[{timestamp}] {message}\n"

        self.log_text.insert(
            tk.END,
            text
        )

        self.log_text.see(
            tk.END
        )

        self.root.update_idletasks()


    # ========================================================
    # UPDATE STATUS
    # ========================================================

    def update_status(self, message):

        self.status_label.config(
            text=message
        )

        self.root.update_idletasks()


    # ========================================================
    # START
    # ========================================================

    def start_scraping(self):

        global STOP_SCRAPING

        STOP_SCRAPING = False


        # ----------------------------------------------------
        # READ SETTINGS
        # ----------------------------------------------------

        prefix = self.prefix_entry.get().strip()

        start_text = self.start_entry.get().strip()

        end_text = self.end_entry.get().strip()

        month_year = self.month_entry.get().strip()

        semester = self.semester_combo.get().strip()


        # ----------------------------------------------------
        # VALIDATE
        # ----------------------------------------------------

        if not prefix:

            messagebox.showerror(
                "Error",
                "Enter PIN prefix."
            )

            return


        try:

            start_pin = int(
                start_text
            )

            end_pin = int(
                end_text
            )

        except:

            messagebox.showerror(
                "Error",
                "Start PIN and End PIN must be numbers."
            )

            return


        if start_pin > end_pin:

            messagebox.showerror(
                "Error",
                "Start PIN cannot be greater than End PIN."
            )

            return


        if not month_year:

            messagebox.showerror(
                "Error",
                "Enter Month / Year."
            )

            return


        if not semester:

            messagebox.showerror(
                "Error",
                "Enter semester."
            )

            return


        # ----------------------------------------------------
        # DISABLE START
        # ----------------------------------------------------

        self.start_button.config(
            state="disabled"
        )


        # ----------------------------------------------------
        # CLEAR LOG
        # ----------------------------------------------------

        self.log_text.delete(
            "1.0",
            tk.END
        )


        # ----------------------------------------------------
        # START THREAD
        # ----------------------------------------------------

        thread = threading.Thread(
            target=self.scrape_thread,
            args=(
                prefix,
                start_pin,
                end_pin,
                month_year,
                semester
            ),
            daemon=True
        )

        thread.start()


    # ========================================================
    # SCRAPE THREAD
    # ========================================================

    def scrape_thread(
        self,
        prefix,
        start_pin,
        end_pin,
        month_year,
        semester
    ):

        global STOP_SCRAPING

        results = []

        total_pins = (
            end_pin - start_pin + 1
        )


        # ----------------------------------------------------
        # START EDGE
        # ----------------------------------------------------

        self.update_status(
            "Starting Microsoft Edge..."
        )

        self.log(
            "Starting Microsoft Edge..."
        )


        try:

            self.driver = create_driver()

            wait = WebDriverWait(
                self.driver,
                15
            )

            self.log(
                "Edge started successfully."
            )

        except Exception as e:

            self.log(
                f"ERROR starting Edge: {e}"
            )

            self.update_status(
                "Could not start Edge."
            )

            self.start_button.config(
                state="normal"
            )

            return


        # ----------------------------------------------------
        # PROCESS PINS
        # ----------------------------------------------------

        try:

            for index, number in enumerate(
                range(
                    start_pin,
                    end_pin + 1
                ),
                start=1
            ):

                if STOP_SCRAPING:

                    self.log(
                        "Scraping stopped by user."
                    )

                    break


                # Always preserve leading zeros
                pin_number = f"{number:03d}"

                pin = (
                    prefix +
                    pin_number
                )


                self.update_status(
                    f"Processing {pin} "
                    f"({index}/{total_pins})"
                )


                self.progress["value"] = (
                    (index - 1) /
                    total_pins *
                    100
                )


                self.log(
                    f"Checking {pin}"
                )


                result = fetch_result(
                    self.driver,
                    wait,
                    pin,
                    month_year,
                    semester
                )


                # ------------------------------------------------
                # SUCCESS
                # ------------------------------------------------

                if result["success"]:

                    result["pin"] = pin

                    results.append(
                        result
                    )


                    self.log(
                        f"SUCCESS: {pin} | "
                        f"{result['student_name']} | "
                        f"Grand Total: "
                        f"{result['grand_total']} | "
                        f"GPA: "
                        f"{result['gpa']} | "
                        f"{result['result']}"
                    )


                    self.log(
                        f"       Papers found: "
                        f"{len(result['papers'])}"
                    )


                # ------------------------------------------------
                # FAILED / NOT FOUND
                # ------------------------------------------------

                else:

                    self.log(
                        f"SKIPPED: {pin} | "
                        f"{result['error']}"
                    )


                self.progress["value"] = (
                    index /
                    total_pins *
                    100
                )


            # ----------------------------------------------------
            # SAVE EXCEL
            # ----------------------------------------------------

            if results:

                timestamp = datetime.now().strftime(
                    "%Y%m%d_%H%M%S"
                )


                filename = (
                    f"SBTET_Results_"
                    f"{semester}_"
                    f"{month_year}_"
                    f"{timestamp}.xlsx"
                )


                output_file = os.path.join(
                    OUTPUT_FOLDER,
                    filename
                )


                self.log(
                    "Creating Excel file..."
                )


                save_excel(
                    results,
                    output_file
                )


                self.log(
                    f"Excel saved successfully:"
                )

                self.log(
                    output_file
                )


                self.update_status(
                    f"Completed - "
                    f"{len(results)} results saved."
                )


                messagebox.showinfo(
                    "Completed",
                    f"Scraping completed.\n\n"
                    f"Students found: {len(results)}\n\n"
                    f"Excel file:\n{output_file}"
                )


            else:

                self.update_status(
                    "No results found."
                )


                messagebox.showwarning(
                    "No Results",
                    "No student results were found."
                )


        except Exception as e:

            self.log(
                f"ERROR: {str(e)}"
            )


            self.update_status(
                "Error occurred."
            )


            messagebox.showerror(
                "Error",
                str(e)
            )


        finally:

            # ------------------------------------------------
            # CLOSE EDGE
            # ------------------------------------------------

            try:

                if self.driver:

                    self.driver.quit()

            except:
                pass


            self.driver = None


            self.start_button.config(
                state="normal"
            )


            self.progress["value"] = 100


    # ========================================================
    # STOP
    # ========================================================

    def stop_scraping(self):

        global STOP_SCRAPING

        STOP_SCRAPING = True

        self.update_status(
            "Stopping..."
        )

        self.log(
            "Stop requested..."
        )


# ============================================================
# MAIN
# ============================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = SBTETDownloader(
        root
    )

    root.mainloop()